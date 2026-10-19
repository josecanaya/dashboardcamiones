# Exporta el deck del comité (deck/project: deck.json + slides/*.html) a PPTX: una imagen 1920×1080 por lámina + notas del orador.
# Uso: python scripts/estado-planta/deck_a_pptx.py <carpeta deck/project> <carpeta con las imágenes /_blob/<id>.png|jpg> <salida.pptx>
# Renderiza cada lámina con Chrome/Edge en modo headless (sin dependencias de Node). Las láminas quedan como imagen (no editables).
import json, os, re, sys, base64, subprocess, tempfile, html as H
from pptx import Presentation
from pptx.util import Emu

deck_dir, blob_dir, out = sys.argv[1:4]
deck = json.load(open(os.path.join(deck_dir, 'deck.json'), encoding='utf-8'))
blobs = {}
for f in os.listdir(blob_dir):
    k, ext = os.path.splitext(f)
    blobs[k] = f'data:{"image/png" if ext == ".png" else "image/jpeg"};base64,' + base64.b64encode(open(os.path.join(blob_dir, f), 'rb').read()).decode()
BROWSERS = [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe']
browser = next(b for b in BROWSERS if os.path.exists(b))
CSS = ('*,*::before,*::after{box-sizing:border-box}html,body{margin:0;padding:0;width:1920px;height:1080px;overflow:hidden;background:#fff}'
       'section{position:relative;width:1920px;height:1080px;overflow:hidden}section>*{position:relative}p,h1,h2,h3,table{margin:0}aside{display:none}img{max-width:none}')
tmp = tempfile.mkdtemp(prefix='deck_pptx_')
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
blank = prs.slide_layouts[6]
for i, sid in enumerate(deck['order']):
    h = open(os.path.join(deck_dir, 'slides', f'{sid}.html'), encoding='utf-8').read()
    m = re.search(r'<aside>(.*?)</aside>', h, re.S); notes = H.unescape(re.sub(r'<[^>]+>', '', m.group(1))).strip() if m else ''
    h = re.sub(r'src="/_blob/([a-f0-9]+)"', lambda mm: f'src="{blobs.get(mm.group(1), "")}"', h)
    page = os.path.join(tmp, f'{i:02d}.html'); png = os.path.join(tmp, f'{i:02d}.png')
    fonts = ''.join(f'<link href="{f["href"]}" rel="stylesheet">' for f in deck.get('faces', {}).values()) or '<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;600;700&display=swap" rel="stylesheet">'
    open(page, 'w', encoding='utf-8').write('<!doctype html><meta charset="utf-8">' + fonts
                                           + f'<style>{CSS}</style>{h}')
    subprocess.run([browser, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--force-device-scale-factor=1', '--virtual-time-budget=4000',
                    '--window-size=1920,1080', f'--screenshot={png}', 'file:///' + page.replace('\\', '/')], check=True, capture_output=True, timeout=120)
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    if notes: s.notes_slide.notes_text_frame.text = notes
    print(f'{i + 1:2d} {sid}')
prs.save(out)
print('PPTX:', out)
