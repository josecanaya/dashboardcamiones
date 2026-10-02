"""Salida editable y extracción de texto heredado; nunca ejecuta generadores antiguos."""
import json, sys, zipfile, textwrap, xml.etree.ElementTree as ET
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

if sys.argv[1] == 'extract':
    with zipfile.ZipFile(sys.argv[2]) as z:
        names = sorted((n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')), key=lambda n: int(Path(n).stem[5:]))
        out = [{'slide': int(Path(n).stem[5:]), 'text': '\n'.join(e.text or '' for e in ET.fromstring(z.read(n)).iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))} for n in names]
    print(json.dumps(out, ensure_ascii=False))
else:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK
    from pptx.oxml.xmlchemy import OxmlElement
    data = json.loads(Path(sys.argv[2]).read_text(encoding='utf-8'))
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    def slide(title, lines, footer='', dark=False):
        s = prs.slides.add_slide(prs.slide_layouts[6]); fill = s.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor.from_string('0B5638' if dark else 'FBFBF8')
        def text(x,y,w,h,string,size,color):
            box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
            for i,line in enumerate(string.split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(size); p.font.name='Arial'; p.font.color.rgb=RGBColor.from_string(color)
        text(.65,.5,12,1.1,title,30,'FFFFFF' if dark else '0B5638')
        text(.65,1.7,12,4.9,'\n'.join(lines),19,'FFFFFF' if dark else '1B2420')
        text(.65,6.8,12,.45,footer[:210],10,'CFE3D7' if dark else '4F5A55')
    def pages(title, lines, footer=''):
        wrapped=[]
        for line in lines:
            wrapped.extend(textwrap.wrap(str(line),width=95) or [''])
        for pos in range(0,len(wrapped),11):
            slide(title+(' · continuación' if pos else ''),wrapped[pos:pos+11],footer)
    def native_chart(section, definition):
        rows=[r for r in definition.get('rows',[]) if len(r)>1 and isinstance(r[1],(int,float))]
        if not rows: return
        hourly='horaria' in definition['title'].lower()
        label=section['title'].replace('ACTIVIDAD EQUIPOS · ','').split(' · ')[0]
        slide(label+' · '+('Actividad horaria' if hourly else 'Operaciones diarias'), [], section.get('source',''))
        s=prs.slides[-1]
        cd=CategoryChartData(); cd.categories=[str(r[0]) for r in rows]; cd.add_series('Conteo de la muestra',[r[1] for r in rows])
        chart=s.shapes.add_chart(XL_CHART_TYPE.LINE if hourly else XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.8),Inches(1.6),Inches(11.7),Inches(4.6),cd).chart
        chart.has_legend=False; chart.chart_style=10
        chart.category_axis.tick_labels.font.size=Pt(10); chart.value_axis.tick_labels.font.size=Pt(11)
        chart.category_axis.major_tick_mark=XL_TICK_MARK.NONE
        chart.value_axis.minimum_scale=0
        chart.value_axis.has_major_gridlines=True
        chart.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string('DDE5DF')
        color=RGBColor.from_string('2E1B4E' if hourly else '0B5638')
        if hourly:
            chart.series[0].format.line.color.rgb=color; chart.series[0].format.line.width=Pt(2)
            skip=OxmlElement('c:tickLblSkip');skip.set('val',str(max(1,len(rows)//8)));chart.category_axis._element.append(skip)
        else:
            chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=color
            chart.series[0].format.line.color.rgb=color
        box=s.shapes.add_textbox(Inches(.8),Inches(6.25),Inches(11.7),Inches(.4))
        p=box.text_frame.paragraphs[0];p.text='Muestra del pipeline; revisar cobertura. '+('Equipo sin atribución de producto.' if section['id'].startswith('equipo') else 'No equivale al universo completo del Excel.')
        p.font.name='Arial';p.font.size=Pt(12);p.font.color.rgb=RGBColor.from_string('4F5A55')
    slide('Comité · Nodo Sur', [f"{data['from']} al {data['to']}", 'Logística y seguridad · versión funcional para revisión', 'Indicadores actuales y antecedentes heredados identificados por separado.'], dark=True)
    for section in data['sections']:
        lines = [section.get('status','')]+[f"{x['label']}: {x['value']} {x.get('unit','')}" for x in section.get('metrics',[])]
        lines += section.get('notes',[])
        for table in section.get('tables',[]):
            rows=[' | '.join(str(v) for v in row) for row in table['rows']]
            lines.extend([table['title'], ' | '.join(table['headers'])]+rows)
        pages(section['title'],lines,section.get('source',''))
        for definition in section.get('charts',[]): native_chart(section,definition)
        for chunk in section.get('referenceSlides',[]):
            text=chunk['text']; pieces=[text[i:i+850] for i in range(0,len(text),850)]
            for part in pieces: pages(f"Antecedente de seguridad · lámina {chunk['slide']}",[part], 'Texto del comité original · no revalidado · '+section.get('source',''))
    prs.save(sys.argv[3])
