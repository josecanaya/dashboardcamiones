from pathlib import Path
from pptx import Presentation

src = Path(r"C:\Users\Usuario\Desktop\Dashboard_camiones\artifacts\claude_import\Reeingeniería logística del Nodo Sur.pptx")
out = src.parent / "inspection" / "contenido.md"
prs = Presentation(src)
lines = [f"# {src.name}", "", f"Slides: {len(prs.slides)}", ""]
for idx, slide in enumerate(prs.slides, 1):
    lines += [f"## Slide {idx}", ""]
    texts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if text:
                texts.append(text)
    lines += texts or ["(sin texto extraíble)"]
    lines.append("")
out.write_text("\n".join(lines), encoding="utf-8")
print(out)
